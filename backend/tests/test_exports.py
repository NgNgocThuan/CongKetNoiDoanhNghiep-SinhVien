from io import BytesIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from tests.school_factories import domain

BASE = "/api/v1/admin/exports"


@pytest.mark.parametrize("kind", ["students", "companies", "jobs", "applications", "skill-report"])
def test_excel_export_permissions_and_content(client, db, make_user, auth_headers, kind):
    d = domain(db, make_user)
    d.student.full_name = "=HYPERLINK(unsafe)"
    db.commit()
    url = f"{BASE}/{kind}.xlsx"
    assert client.get(url, headers=auth_headers(d.student_user)).status_code == 403
    result = client.get(url, headers=auth_headers(d.admin))
    assert result.status_code == 200, result.text
    assert "spreadsheetml" in result.headers["content-type"]
    assert "attachment" in result.headers["content-disposition"]
    sheet = load_workbook(BytesIO(result.content)).active
    assert sheet.max_row >= 1
    if kind == "students":
        assert sheet.cell(2, 2).value == "'=HYPERLINK(unsafe)"
        assert sheet.cell(2, 5).value is None
        text = str(list(sheet.values))
        assert "0901234567" not in text and "private@test.com" not in text
    assert (
        client.get(
            url + "?date_from=2027-01-01&date_to=2026-01-01", headers=auth_headers(d.admin)
        ).status_code
        == 422
    )


def test_pdf_aggregate_report(client, db, make_user, auth_headers):
    d = domain(db, make_user)
    url = f"{BASE}/dashboard.pdf"
    assert client.get(url, headers=auth_headers(d.company_user)).status_code == 403
    result = client.get(url, headers=auth_headers(d.admin))
    assert result.status_code == 200
    assert result.content.startswith(b"%PDF")
    pdf = PdfReader(BytesIO(result.content))
    text = "".join(page.extract_text() for page in pdf.pages)
    assert "Student Connect" in text and "Applications" in text
    assert d.student.full_name not in text
    assert (
        client.get(
            url + "?date_from=2027-01-01&date_to=2026-01-01", headers=auth_headers(d.admin)
        ).status_code
        == 422
    )


def test_empty_pdf_report(client, make_user, auth_headers):
    from app.core.enums import UserRole

    response = client.get(f"{BASE}/dashboard.pdf", headers=auth_headers(make_user(UserRole.ADMIN)))
    assert response.status_code == 200 and response.content.startswith(b"%PDF")
