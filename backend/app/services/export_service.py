"""Export projections deliberately omit contact details, birth dates and CV text."""

from io import BytesIO

from openpyxl import Workbook
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.exceptions import NotFoundError
from app.models import Application, Company, Job, Major, Student
from app.schemas.common import PageParams
from app.services import dashboard_service as dashboard


def _safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _workbook(headers, rows):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Report")
    sheet.append(headers)
    for row in rows:
        sheet.append([_safe_cell(v) for v in row])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_excel(db: Session, kind, filters):
    if kind == "students":
        # GPA is included only with explicit permission, even for bulk exports.
        from sqlalchemy import case

        from app.models import StudentProfile

        stmt = (
            select(
                Student.student_code,
                Student.full_name,
                Major.name,
                Student.cohort,
                case((StudentProfile.show_gpa.is_(True), Student.gpa), else_=None),
            )
            .join(Major)
            .outerjoin(StudentProfile)
            .where(*dashboard.student_conditions(filters))
            .order_by(Student.id)
        )
        headers = ["Student code", "Full name", "Major", "Cohort", "GPA (consented)"]
    elif kind == "companies":
        stmt = (
            select(Company.id, Company.name, Company.verification_status)
            .where(*dashboard.time_conditions(Company.created_at, filters))
            .order_by(Company.id)
        )
        if filters.term_id or filters.major_id or filters.faculty_id:
            stmt = stmt.where(
                Company.id.in_(
                    select(Job.company_id).where(*dashboard.job_conditions(filters, dates=False))
                )
            )
        headers = ["ID", "Company", "Verification"]
    elif kind == "jobs":
        stmt = (
            select(Job.id, Job.title, Job.company_id, Job.term_id, Job.status, Job.deadline)
            .where(*dashboard.job_conditions(filters))
            .order_by(Job.id)
        )
        headers = ["ID", "Title", "Company ID", "Term ID", "Status", "Deadline"]
    elif kind == "applications":
        stmt = (
            select(
                Application.id,
                Student.student_code,
                Application.job_id,
                Application.current_status,
                Application.applied_at,
            )
            .join(Student)
            .join(Job)
            .where(*dashboard.application_conditions(filters))
            .order_by(Application.id)
        )
        headers = ["ID", "Student code", "Job ID", "Status", "Applied at"]
    elif kind == "skill-report":

        def skill_rows():
            page = 1
            while True:
                result = dashboard.skill_gap(db, filters, PageParams(page=page, size=100))
                for item in result["items"]:
                    yield [
                        item[k]
                        for k in (
                            "name",
                            "required_jobs",
                            "students_with_skill",
                            "students_missing_skill",
                        )
                    ]
                if page >= result["pages"]:
                    break
                page += 1

        return _workbook(
            ["Skill", "Jobs requesting", "Students with skill", "Students missing skill"],
            skill_rows(),
        )
    else:
        raise NotFoundError("Không tìm thấy loại báo cáo")
    return _workbook(headers, db.execute(stmt.execution_options(yield_per=500)))


def export_dashboard_pdf(db: Session, filters):
    stats = dashboard.overview(db, filters)
    output = BytesIO()
    doc = SimpleDocTemplate(output)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Student Connect - Dashboard", styles["Title"]),
        Paragraph("Aggregate statistics. No individual student data.", styles["Normal"]),
        Spacer(1, 14),
    ]
    if filters.date_from or filters.date_to:
        story.append(
            Paragraph(
                f"Period: {filters.date_from or 'start'} to {filters.date_to or 'present'}",
                styles["Normal"],
            )
        )
    labels = {
        "students": "Students",
        "companies": "Companies",
        "jobs": "Jobs",
        "applications": "Applications",
        "interviews": "Current interviews",
        "hired": "Hired",
    }
    story.append(Table([["Metric", "Count"]] + [[labels[k], str(v)] for k, v in stats.items()]))
    drawing = Drawing(430, 210)
    chart = VerticalBarChart()
    chart.x, chart.y, chart.width, chart.height = 35, 45, 370, 140
    chart.data = [list(stats.values())]
    chart.categoryAxis.categoryNames = list(labels.values())
    chart.categoryAxis.labels.angle = 25
    chart.categoryAxis.labels.fontSize = 7
    chart.valueAxis.valueMin = 0
    chart.bars[0].fillColor = colors.HexColor("#2563eb")
    drawing.add(chart)
    story.extend([Spacer(1, 20), drawing])
    doc.build(story)
    return output.getvalue()
