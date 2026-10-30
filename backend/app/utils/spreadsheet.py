"""Excel/CSV readers with stable physical row numbers and business errors."""

import csv
import io
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.exceptions import BadRequestError


def read_numbered_rows(filename: str, content: bytes) -> list[tuple[int, dict[str, str]]]:
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".csv":
            data = list(csv.reader(io.StringIO(content.decode("utf-8-sig")), strict=True))
        elif ext == ".xlsx":
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                data = list(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
        else:
            raise BadRequestError("Chỉ hỗ trợ file .xlsx hoặc .csv")
    except (
        UnicodeError,
        csv.Error,
        BadZipFile,
        InvalidFileException,
        ValueError,
        KeyError,
        ParseError,
        OSError,
        EOFError,
    ) as exc:
        raise BadRequestError("File Excel/CSV không hợp lệ hoặc không dùng mã hóa UTF-8") from exc
    if not data:
        raise BadRequestError("File không có tiêu đề")
    headers = [str(h).strip() if h is not None else "" for h in data[0]]
    if not headers or any(not h for h in headers) or len(set(headers)) != len(headers):
        raise BadRequestError("Tiêu đề cột trống hoặc bị trùng")
    cleaned = []
    for number, values in enumerate(data[1:], 2):
        if not any(v is not None and str(v).strip() for v in values):
            continue
        row = {
            key: "" if i >= len(values) or values[i] is None else str(values[i]).strip()
            for i, key in enumerate(headers)
        }
        if len(values) > len(headers) and any(
            v is not None and str(v).strip() for v in values[len(headers) :]
        ):
            row["_import_error"] = "Dòng có nhiều ô hơn tiêu đề"
        cleaned.append((number, row))
    return cleaned


def read_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    """Keep the shared API: first row is headers; blank rows are skipped."""
    return [row for _, row in read_numbered_rows(filename, content)]
